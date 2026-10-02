from docx import Document
from pathlib import Path
import shutil

base = Path(r'C:\Users\Israel\Documents\projetos\campanha_do_agasalho')
source = base / 'apostila_campanha_do_agasalho.docx'

if source.exists():
    doc = Document(str(source))
else:
    doc = Document()
    doc.add_heading('APOSTILA DE PROJETO', 0)
    doc.add_heading('Campanha do Agasalho', level=1)
    doc.add_paragraph('Do formulário ao Docker, construindo passo a passo')
    doc.add_paragraph('6 aulas de 50 minutos | Carga horária: 5 horas')
    doc.add_paragraph('Node.js • Express • SQLite • HTML/CSS/Bootstrap • Docker')
    doc.add_paragraph('Material do professor e guia de prática do estudante')

for _ in range(3):
    doc.add_paragraph('')
doc.add_page_break()
doc.add_heading('MATERIAL COMPLEMENTAR: VISÃO GERAL DIDÁTICA DO FLUXO', level=1)

sections = [
    ('1) Estrutura do HTML em public/index.html', [
        'A página HTML é a camada de apresentação da aplicação. Ela define o que o usuário vê e onde os formulários, botões, modais e listas são montados.',
        'A estrutura principal contém elementos como header, hero, formulário de doação, área administrativa, modal de login e modal de agradecimento.',
        'O navegador lê esse arquivo e monta a interface visual. O JavaScript, por sua vez, busca esses elementos no DOM para adicionar comportamento com eventos e requisições.',
        'No formulário público, o estudante digita nome, e-mail, telefone, cidade, tamanho e observação. Esses campos são enviados ao backend em formato de objeto JavaScript.',
        'A área administrativa fica escondida até o administrador entrar. Isso é uma boa prática para limitar o acesso e separar a interface pública da interface privada.'
    ]),
    ('2) O que acontece em public/app.js', [
        'Este arquivo é o cérebro do frontend. Ele seleciona elementos do HTML, escuta cliques, lê formulários e conversa com o servidor por meio de fetch.',
        "A função chamarApi centraliza o acesso à API. Ela recebe a URL, as opções da requisição e trata erros de rede e do servidor.",
        "Exemplo de padrão: const resposta = await fetch(url, { method: 'POST', headers: { 'Content-Type': 'application/json' }, body: JSON.stringify(dados) });",
        'A função mostrarDoadores monta a lista visual com dados vindos do backend. O uso de map permite transformar cada doador em uma linha de HTML, e o join junta todas as linhas em uma única string.',
        'A função carregarDoadores chama a API, recebe os dados e renderiza a tabela ou cards na tela. Isso acontece sem recarregar a página.',
        'O código também cuida de mensagens de erro e sucesso para que a experiência do usuário seja clara e confiável.'
    ]),
    ('3) Fluxo do cadastro de doador', [
        'Ao enviar o formulário público, o JavaScript captura os dados com Object.fromEntries(new FormData(formularioDoador).entries()).',
        'Esse comando transforma todos os campos do formulário em um objeto JavaScript simples, como { nome: ..., email: ..., cidade: ... }.',
        'Se existir um doador em edição, a aplicação envia uma requisição de atualização. Caso contrário, envia uma requisição de criação.',
        'A chamada POST ou PUT é enviada para /api/doadores. O backend valida os dados, salva no banco e responde com sucesso ou erro.',
        'Após a resposta positiva, o app limpa o formulário, atualiza a listagem e abre o modal de agradecimento, fechando o ciclo da interação.'
    ]),
    ('4) Fluxo de edição e exclusão', [
        'O botão editar extrai o identificador do doador e busca os dados no backend. Em seguida, preenche os campos do formulário com os valores atuais.',
        'Quando o aluno clica em excluir, o código pede confirmação com confirm(). Isso evita exclusão acidental de dados importantes.',
        'A partir daí, a API recebe uma requisição DELETE e apaga o registro referente ao ID informado.',
        'A lista da tela é recarregada automaticamente para refletir a mudança, reforçando a ideia de que a interface é resultado da comunicação contínua com o servidor.'
    ]),
    ('5) Fluxo de login e administração', [
        'O login administrativo usa um formulário próprio. O código envia usuário e senha para /api/entrar.',
        'O backend compara a senha informada com o hash armazenado no banco. Isso é um conceito de segurança essencial: o sistema não salva a senha em texto puro.',
        'Se a autenticação for válida, o servidor cria uma sessão e envia um cookie para o navegador.',
        'A partir do momento em que o cookie é enviado, o navegador inclui essa informação nas próximas requisições. O backend valida a sessão em cada rota protegida.',
        'Ao sair, a sessão é removida e o cookie é expirado. A interface administrativa some e o usuário retorna à experiência pública.'
    ]),
    ('6) O que acontece em server.js', [
        'O arquivo server.js é o núcleo da aplicação. Ele configura o Express, cria a conexão com o banco, define as rotas e cuida da autenticação.',
        'O projeto foi preparado para funcionar com SQLite em desenvolvimento e PostgreSQL em produção. Isso fica visível por meio da detecção de DATABASE_URL.',
        'As funções criarBancoSqlite e criarBancoPostgres implementam essa diferença de ambiente. Em desenvolvimento, o sistema usa o arquivo local em data/doadores.sqlite.',
        'A criação das tabelas usa CREATE TABLE IF NOT EXISTS para evitar erros caso o banco já exista.',
        'A função exigirAdministrador atua como middleware. Ela verifica se a sessão está ativa antes de permitir acesso a uma rota sensível.'
    ]),
    ('7) Rotas principais', [
        'POST /api/entrar: autentica usuário e senha e cria sessão.',
        'POST /api/sair: encerra a sessão e limpa o cookie.',
        'GET /api/doadores: lista os doadores e pode receber busca por texto.',
        'POST /api/doadores: cria um novo cadastro público.',
        'PUT /api/doadores/:id: atualiza um registro existente.',
        'DELETE /api/doadores/:id: remove a linha do banco.',
        'As rotas administrativas normalmente usam exigirAdministrador, garantindo que apenas usuários autenticados possam acessá-las.'
    ]),
    ('8) Como explicar em sala de aula', [
        'Para facilitar a compreensão, a aula pode seguir um roteiro simples: entrada, validação, banco, resposta e atualização visual.',
        'No início, o professor pode mostrar o fluxo completo em uma linha: formulário → JavaScript → fetch → backend → banco → resposta → interface.',
        'Depois, é útil detalhar cada camada separadamente: HTML define o formulário; JavaScript coleta os dados; Express recebe a requisição; SQLite ou PostgreSQL grava a informação; o cliente atualiza o resultado.',
        'Esse formato ajuda os estudantes a enxergar a arquitetura sem se perder nos detalhes de implementação.'
    ]),
    ('9) Resumo do raciocínio', [
        'A aplicação funciona como uma cadeia de comunicação entre camadas: interface, lógica do cliente, API e banco de dados.',
        'O HTML prepara a tela; o JavaScript transforma ações em requisições; o Node.js/Express valida e processa dados; o banco persistente armazena informações; e a interface atualiza o estado visual da aplicação.',
        'Em outras palavras, o projeto ensina, na prática, como uma web app funciona em arquitetura full-stack, com autenticação, persistência e fluxo de dados real.'
    ]),
]

for title, lines in sections:
    doc.add_heading(title, level=2)
    for line in lines:
        doc.add_paragraph(line)

doc.add_heading('Explicação linha por linha e comandos importantes', level=1)
for line in [
    'A seguir, alguns trechos do projeto são comentados em linguagem didática para facilitar o ensino.',
    "1. Coleta de formulário: Object.fromEntries(new FormData(formularioDoador).entries())",
    'Esse comando transforma os campos do formulário em um objeto JavaScript. É a base para enviar dados estruturados ao backend.',
    '2. Renderização com map: listaDoadores.innerHTML = doadores.map(...).join(\"\");',
    'O método map percorre cada doador e cria uma linha visual por item. Em seguida, join junta todas as linhas em uma string única que é inserida no HTML.',
    "3. Chamada de API: fetch('/api/doadores', { method: 'POST', headers: { 'Content-Type': 'application/json' }, body: JSON.stringify(dados) })",
    'Essa linha representa a comunicação entre frontend e backend. O cliente envia um JSON contendo os dados do formulário.',
    '4. Validação de resposta: if (!resposta.ok) throw new Error(...)',
    'Se o servidor responder com erro, o código interrompe o fluxo e exibe a mensagem correta ao usuário.',
    "5. Inserção no banco: db.prepare('INSERT INTO doadores (...) VALUES (?, ?, ...)').run(...)",
    'A instrução SQL usa placeholders (?, ?, ...) para evitar concatenação de texto. Isso reduz risco de quebra e melhora a segurança.',
    '6. Middleware de sessão: if (!req.session || !req.session.adminId) return res.status(401).json(...)',
    'Esse trecho controla quem pode acessar áreas privadas da aplicação. Se o usuário não estiver autenticado, a API rejeita a requisição.',
    '7. Hash da senha: scryptSync(senha, salt, 64)',
    'A senha não é guardada em texto puro. Em vez disso, o servidor gera um hash que só pode ser comparado por meio de uma validação específica.',
    '8. Filtro de busca: doadores.filter(doador => doador.cidade.toLowerCase().includes(busca.toLowerCase()))',
    'Esse filtro permite procurar registros dentro da listagem sem recarregar a página.'
]:
    doc.add_paragraph(line)

doc.add_heading('Material para o professor', level=1)
for line in [
    'Objetivo pedagógico: mostrar aos alunos como uma aplicação web full-stack organiza interface, API, autenticação e persistência.',
    'Roteiro da apresentação:',
    '1. Abrir o projeto e mostrar a estrutura básica: public/, api/, data/, package.json e server.js.',
    '2. Explicar a diferença entre frontend e backend com uma analogia simples: a interface faz pedidos e o servidor executa a regra de negócio.',
    '3. Demonstrar o formulário de doação e mostrar como os dados são transformados em objeto.',
    '4. Explicar o fluxo de fetch e a resposta do servidor.',
    '5. Mostrar a criação das tabelas e a operação de INSERT.',
    '6. Apresentar a autenticação com sessão e cookie.',
    '7. Mostrar busca, edição e exclusão no painel administrativo.',
    '8. Encerrar com a visão geral do ciclo: navegação → ação → API → banco → resposta → interface.',
    'Sugestão de fala do professor: “Hoje vamos entender, com um projeto real, como uma aplicação web transforma uma ação do usuário em uma operação no banco e em uma resposta visual.”',
    'Dica de avaliação: peça que o aluno descreva, em duas frases, o que acontece quando o usuário clica em cadastrar doador.',
    'Dica de atividade: peça para cada dupla explicar a responsabilidade de uma camada: HTML, app.js, server.js e banco.'
]:
    doc.add_paragraph(line)

doc.add_heading('Versão para alunos – nível básico', level=1)
for line in [
    'A camada HTML é a tela. A camada JavaScript faz o site reagir. O backend recebe os dados e guarda no banco.',
    'Quando o usuário clica em “Cadastrar”, o navegador monta um objeto com os dados. Depois, o JavaScript envia esse objeto para o servidor.',
    'O servidor valida. Se tudo estiver correto, salva no banco. Então envia uma confirmação para a tela.',
    'O mesmo padrão vale para login, busca e edição. Em todas as operações, a lógica geral é: usuário → ação → requisição → processamento → resposta.',
    'Se o aluno entender esse ciclo, ele já compreende grande parte da arquitetura do projeto.',
    'O conceito de map é simples: ele percorre a lista e transforma cada item em um pedaço de HTML. O join então monta tudo em um único bloco.',
    'O conceito de sessão é simples: ao fazer login, o servidor cria um identificador e guarda a informação. Nas próximas requisições, o navegador entrega esse identificador e o servidor reconhece o usuário.'
]:
    doc.add_paragraph(line)

doc.add_heading('Versão para alunos – nível avançado', level=1)
for line in [
    'A aplicação funciona em camadas com responsabilidade definida: interface, lógica de cliente, API, autenticação, persistência e resposta.',
    'No frontend, a função Object.fromEntries(new FormData(formulario).entries()) é usada para transformar entrada de formulário em objeto. Essa etapa evita escrever manualmente cada atributo.',
    'No backend, Express recebe os dados em req.body e valida campos básicos. A operação SQL é executada com placeholders para evitar SQL injection e garantir integridade.',
    "Uso típico de map em JavaScript: doadores.map(doador => '...').join(''). Isso produz HTML dinâmico para a lista de registros.",
    'O middleware exigirAdministrador é o ponto de decisão para rotas protegidas. Ele verifica se a sessão existe e se ela corresponde a um administrador válido antes de prosseguir.',
    'A operação de hash com scryptSync reforça segurança. O sistema compara a entrada informada com o hash guardado no banco, não com a senha em texto puro.',
    'A arquitetura também contempla multi ambiente: SQLite para desenvolvimento e PostgreSQL para produção via DATABASE_URL. Isso permite o mesmo código rodar em ambientes distintos com pequenos ajustes de configuração.',
    'A vantagem didática dessa solução é que cada camada pode ser estudada isoladamente, mas o aluno consegue ver o funcionamento completo do fluxo.',
    'Em resumo, o projeto ensina profissionalmente como uma aplicação web segue princípios de UX, segurança, persistência e arquitetura de software.'
]:
    doc.add_paragraph(line)

source_path = str(source)
doc.save(source_path)

print('Documento criado com sucesso:', source_path)
